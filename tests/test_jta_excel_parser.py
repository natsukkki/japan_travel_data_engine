"""Tests for the Excel-only JTA parser; no MinIO or database access."""

from io import BytesIO
import unittest

from openpyxl import Workbook

from processing.spark.jta_excel_parser import parse_jta_workbook


def workbook_bytes(months: list[int], foreign_column: int = 17) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for month in months:
        sheet = workbook.create_sheet(f"第2表({month}月)")
        sheet.cell(4, 2, "延べ\n宿泊者数\n1)")
        sheet.cell(4, foreign_column, "うち\n外国人延べ\n宿泊者数\n1)")
        for code in range(1, 48):
            row = code + 7
            sheet.cell(row, 1, f" {code:02d}県{code}")
            sheet.cell(row, 2, 1000 + code)
            sheet.cell(row, foreign_column, 100 + code)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class JtaExcelParserTests(unittest.TestCase):
    def test_preliminary_reads_prefectures_and_dynamic_foreign_column(self) -> None:
        rows = parse_jta_workbook(
            workbook_bytes([6], foreign_column=20),
            release_type="second_preliminary", year=2026, month=6,
        )
        self.assertEqual(len(rows), 47)
        self.assertEqual(rows[0], {
            "prefecture_code": 1, "prefecture_name_jp": "県1",
            "year": 2026, "month": 6,
            "total_guest_nights": 1001, "foreign_guest_nights": 101,
        })

    def test_final_reads_twelve_months_without_annual_total(self) -> None:
        rows = parse_jta_workbook(
            workbook_bytes(list(range(1, 13))), release_type="final", year=2025,
        )
        self.assertEqual(len(rows), 12 * 47)
        self.assertEqual({row["month"] for row in rows}, set(range(1, 13)))

    def test_rejects_missing_month(self) -> None:
        with self.assertRaisesRegex(ValueError, "Неверный набор"):
            parse_jta_workbook(workbook_bytes([1]), release_type="final", year=2025)

    def test_rejects_foreign_count_above_total(self) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet = workbook.create_sheet("第2表(6月)")
        sheet.cell(4, 2, "延べ宿泊者数")
        sheet.cell(4, 17, "外国人延べ宿泊者数")
        sheet.cell(8, 1, "01北海道")
        sheet.cell(8, 2, 10)
        sheet.cell(8, 17, 11)
        output = BytesIO()
        workbook.save(output)
        with self.assertRaisesRegex(ValueError, "больше общего"):
            parse_jta_workbook(output.getvalue(), release_type="second_preliminary", year=2026, month=6)


if __name__ == "__main__":
    unittest.main()
